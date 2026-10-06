#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""tw_calendar.py — 台股交易日曆 (唯一真相來源)。

為什麼要這支 (2026-09-28 用戶:「9/8 T-3 但是今天休市欸」):
  訂價窗口的 T-5 / T-3 / T-1 是「訂價基準日前 1/3/5 個【營業日】的收盤價」,
  但原本 HTML 的 pwBizAdd/pwBizDiff 只跳週末、不跳國定假日:
      2026-09-25 中秋 + 09-28 教師節補休 連在一起 →
      T=10/01 時,T-3 被算成 09/28 (休市,根本沒有收盤價),正確是 09/24。
  Python 端也有兩處同樣寫法 (cb_listing_advisor.next_trading_day、run_update 取前一交易日)。
  → 全部改用本模組,假日表集中一份。

資料來源:統一證 xlsx 的 `holiday` 分頁 (他們自己維護,含颱風假,已到 2027-12-31)。
  每次 backfill_primary_dates 跑就同步進 DB 的 holidays 表 → 本模組與 build_html 都讀 DB,
  就算某天 xlsx 沒進來也還有上一份。

⚠ 覆蓋範圍會過期:假日表只到 2027-12-31,超出範圍就退化成「只跳週末」(會錯)。
  `coverage_ok()` 供稽核呼叫,快到期要提醒補表。
"""
import sqlite3
from datetime import date, datetime, timedelta
from pathlib import Path

HERE = Path(__file__).parent
DB_PATH = HERE / 'cb_data.db'
CB_DROP = Path(r'C:\Users\J.Chun\Desktop\stock-dash\cbas-template\CB報')

_CACHE = None      # set[str] of 'YYYY-MM-DD'
_CACHE_RANGE = None


def ensure_table(conn):
    conn.execute('''CREATE TABLE IF NOT EXISTS holidays (
        date TEXT PRIMARY KEY, name TEXT, source TEXT, updated_at TEXT)''')


def import_from_xlsx(conn, path=None, verbose=False):
    """把統一證 xlsx 的 holiday 分頁同步進 DB。回 (新增, 總筆數)。"""
    import glob
    import openpyxl
    if path is None:
        files = sorted(glob.glob(str(CB_DROP / 'CB發行資訊與CBAS報價表_統一證_*.xlsx')))
        if not files:
            return 0, 0
        path = files[-1]
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if 'holiday' not in wb.sheetnames:
            wb.close()
            return 0, 0
        ws = wb['holiday']
        rows = []
        for r in ws.iter_rows(values_only=True):
            if not r or r[0] is None:
                continue
            v = r[0]
            iso = v.strftime('%Y-%m-%d') if isinstance(v, (datetime, date)) else str(v)[:10]
            if len(iso) == 10 and iso[4] == '-':
                rows.append((iso, str(r[1] or '').strip()))
        wb.close()
    except Exception as e:
        if verbose:
            print(f'  [WARN] holiday 分頁讀取失敗: {e}')
        return 0, 0
    ensure_table(conn)
    before = conn.execute('SELECT COUNT(*) FROM holidays').fetchone()[0]
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn.executemany(
        'INSERT INTO holidays(date,name,source,updated_at) VALUES(?,?,?,?) '
        'ON CONFLICT(date) DO UPDATE SET name=excluded.name, updated_at=excluded.updated_at',
        [(d, n, 'unisec-xlsx', now) for d, n in rows])
    after = conn.execute('SELECT COUNT(*) FROM holidays').fetchone()[0]
    global _CACHE, _CACHE_RANGE
    _CACHE = _CACHE_RANGE = None      # 讓下次讀重新載入
    return after - before, after


def holidays(conn=None):
    """回 set('YYYY-MM-DD')。第一次讀之後快取在記憶體。"""
    global _CACHE, _CACHE_RANGE
    if _CACHE is not None:
        return _CACHE
    own = conn is None
    if own:
        if not DB_PATH.exists():
            _CACHE, _CACHE_RANGE = set(), None
            return _CACHE
        conn = sqlite3.connect(str(DB_PATH))
    try:
        ensure_table(conn)
        rows = [r[0] for r in conn.execute('SELECT date FROM holidays')]
    except Exception:
        rows = []
    finally:
        if own:
            conn.close()
    _CACHE = set(rows)
    _CACHE_RANGE = (min(rows), max(rows)) if rows else None
    return _CACHE


def coverage():
    """回 (最早, 最晚) 假日日期;空表回 None。"""
    holidays()
    return _CACHE_RANGE


def coverage_ok(through=None):
    """假日表是否涵蓋到指定日期 (預設:今天 +120 天)。稽核用。"""
    rng = coverage()
    if not rng:
        return False
    end = through or (date.today() + timedelta(days=120)).isoformat()
    return rng[1] >= end


def _d(x):
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    return date.fromisoformat(str(x)[:10])


def is_trading_day(d, hols=None):
    d = _d(d)
    hols = holidays() if hols is None else hols
    return d.weekday() < 5 and d.isoformat() not in hols


def biz_add(d, n, hols=None):
    """往後(正)/往前(負)推 n 個交易日。n=0 回原日 (不校正成交易日)。"""
    d = _d(d)
    if n == 0:
        return d
    hols = holidays() if hols is None else hols
    step = 1 if n > 0 else -1
    cnt = 0
    guard = 0
    while cnt != n:
        d += timedelta(days=step)
        guard += 1
        if guard > 3000:
            raise RuntimeError('biz_add 沒收斂 — 假日表可能異常')
        if is_trading_day(d, hols):
            cnt += step
    return d


def biz_diff(a, b, hols=None):
    """a → b 之間有幾個交易日 (b 在 a 之後為正)。起點不計、終點計。"""
    a, b = _d(a), _d(b)
    if a == b:
        return 0
    hols = holidays() if hols is None else hols
    step = 1 if b > a else -1
    n = 0
    d = a
    guard = 0
    while d != b:
        d += timedelta(days=step)
        guard += 1
        if guard > 3000:
            raise RuntimeError('biz_diff 沒收斂')
        if is_trading_day(d, hols):
            n += step
    return n


def next_trading_day(d=None, hols=None):
    d = _d(d or date.today())
    hols = holidays() if hols is None else hols
    d += timedelta(days=1)
    while not is_trading_day(d, hols):
        d += timedelta(days=1)
    return d


def prev_trading_day(d=None, hols=None):
    d = _d(d or date.today())
    hols = holidays() if hols is None else hols
    d -= timedelta(days=1)
    while not is_trading_day(d, hols):
        d -= timedelta(days=1)
    return d


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser(description='台股交易日曆')
    ap.add_argument('--import-xlsx', action='store_true', help='從統一證 xlsx 同步假日表進 DB')
    ap.add_argument('--check', help='檢查某天是不是交易日 (YYYY-MM-DD)')
    args = ap.parse_args()
    conn = sqlite3.connect(str(DB_PATH))
    if args.import_xlsx:
        added, total = import_from_xlsx(conn, verbose=True)
        conn.commit()
        print(f'假日表同步: 新增 {added} / 共 {total} 筆')
    rng = coverage()
    print(f'涵蓋範圍: {rng[0]} ~ {rng[1]}' if rng else '假日表是空的')
    print(f'涵蓋到今天+120天? {"是" if coverage_ok() else "否 — 該補表了"}')
    if args.check:
        d = date.fromisoformat(args.check)
        print(f'{d} → {"交易日" if is_trading_day(d) else "休市"}')
    conn.close()
